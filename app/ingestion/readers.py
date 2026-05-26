"""
File readers for PDF, Word (.docx), and image files.

PDF pages are extracted natively via PyMuPDF. Pages that contain fewer
characters than the configured threshold are treated as scanned images and
re-processed with Tesseract OCR. Standalone image files are also run through
Tesseract. Word documents are read via python-docx.

Each reader returns a list of page dicts with the shape:
    {
        "page":   int,       # 1-based page number
        "text":   str,       # raw extracted text
        "urls":   list[str], # URLs found in this page
        "source": str,       # original file path
        "method": str,       # "native" | "ocr"
    }
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List

import fitz  # PyMuPDF
import pytesseract
from docx import Document
from PIL import Image, ImageFilter, ImageOps

from app.config import settings

pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD

SUPPORTED_EXTENSIONS: set[str] = {
    ".pdf", ".docx", ".doc",
    ".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp",
}

_URL_RE = re.compile(
    r"https?://[^\s‌‍‎‏‪-‮]+"
    r"|www\.[^\s‌‍‎‏‪-‮]+"
)


def _extract_urls(text: str) -> List[str]:
    return _URL_RE.findall(text)


def _preprocess_for_ocr(image: Image.Image) -> Image.Image:
    """Apply contrast, denoise, sharpen, and binarise before OCR."""
    image = image.convert("L")
    image = ImageOps.autocontrast(image)
    image = image.filter(ImageFilter.MedianFilter(size=3))
    image = image.filter(ImageFilter.SHARPEN)
    image = image.point(lambda px: 0 if px < 160 else 255, "1")
    return image


def _ocr(image: Image.Image) -> str:
    processed = _preprocess_for_ocr(image)
    return pytesseract.image_to_string(processed, lang=settings.OCR_LANG)


# ── Public readers ────────────────────────────────────────────────────────────

def read_pdf(path: Path) -> List[Dict[str, Any]]:
    """Extract text from every page of a PDF file.

    Falls back to Tesseract OCR for pages whose native text extraction
    yields fewer than OCR_FALLBACK_THRESHOLD characters (scanned pages).
    """
    doc = fitz.open(str(path))
    pages: List[Dict[str, Any]] = []

    for i, page in enumerate(doc):
        native_text = page.get_text("text")

        if len(native_text.strip()) >= settings.OCR_FALLBACK_THRESHOLD:
            text = native_text
            method = "native"
        else:
            pix = page.get_pixmap(dpi=300)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text = _ocr(img)
            method = "ocr"

        pages.append({
            "page":   i + 1,
            "text":   text,
            "urls":   _extract_urls(text),
            "source": str(path),
            "method": method,
        })

    doc.close()
    return pages


def read_word(path: Path) -> List[Dict[str, Any]]:
    """Extract all paragraphs and tables from a Word document."""
    doc = Document(str(path))
    parts: List[str] = []

    for para in doc.paragraphs:
        stripped = para.text.strip()
        if stripped:
            parts.append(stripped)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts)
    return [{
        "page":   1,
        "text":   full_text,
        "urls":   _extract_urls(full_text),
        "source": str(path),
        "method": "native",
    }]


def read_image(path: Path) -> List[Dict[str, Any]]:
    """Run Tesseract OCR on a standalone image file."""
    image = Image.open(str(path))
    text = _ocr(image)
    return [{
        "page":   1,
        "text":   text,
        "urls":   _extract_urls(text),
        "source": str(path),
        "method": "ocr",
    }]


def read_file(path: Path) -> List[Dict[str, Any]]:
    """Dispatch to the correct reader based on file extension."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext in {".docx", ".doc"}:
        return read_word(path)
    if ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"}:
        return read_image(path)
    raise ValueError(f"Unsupported file type: '{ext}'")
